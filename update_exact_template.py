import sys
import os
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

template_path = 'C:/Users/shine.DESKTOP-HU962IE/.gemini/antigravity/brain/3bc3d249-8bc9-4793-bc90-80e80f426408/.user_uploaded/media_1791993717891_e6c58a20.pptx'
output_path = 'C:/Users/shine.DESKTOP-HU962IE/.gemini/antigravity/scratch/statskill-ai/SIH_PPT_AIZENIX_21121_Template.pptx'
alt_output_path = 'C:/Users/shine.DESKTOP-HU962IE/.gemini/antigravity/scratch/statskill-ai/SIH_PPT-1.pptx'

prs = Presentation(template_path)

# =========================================================================
# SLIDE 0: Title Slide
# =========================================================================
s0 = prs.slides[0]
tf0 = s0.shapes[3].text_frame

# Keep formatting clean
# p[1]: Problem Statement Title
p1 = tf0.paragraphs[1]
p1.clear()
r1_label = p1.add_run()
r1_label.text = "Problem Statement Title - "
r1_label.font.bold = True
r1_label.font.size = Pt(20)
r1_val = p1.add_run()
r1_val.text = "StatSkill AI: AI-Assisted Skill Assessment & Competency Calibration Platform"
r1_val.font.bold = False
r1_val.font.size = Pt(20)

# p[2]: Insert Idea / PS ID line or replace
# Let's inspect paragraphs in tf0
# Currently: p[0]='', p[1]=Title, p[2]=Theme, p[3]=Category, p[4]=Team ID, p[5]=Team Name
p2 = tf0.paragraphs[2]
p2.clear()
r2_label = p2.add_run()
r2_label.text = "Problem Statement / Idea ID - "
r2_label.font.bold = True
r2_label.font.size = Pt(22)
r2_val = p2.add_run()
r2_val.text = "21121"
r2_val.font.bold = False
r2_val.font.size = Pt(22)

p3 = tf0.paragraphs[3]
p3.clear()
r3_label = p3.add_run()
r3_label.text = "Theme - "
r3_label.font.bold = True
r3_label.font.size = Pt(22)
r3_val = p3.add_run()
r3_val.text = "Smart Education    |    PS Category - Software"
r3_val.font.bold = False
r3_val.font.size = Pt(22)

p4 = tf0.paragraphs[4]
p4.clear()
r4_label = p4.add_run()
r4_label.text = "Team ID - "
r4_label.font.bold = True
r4_label.font.size = Pt(22)
r4_val = p4.add_run()
r4_val.text = "176704"
r4_val.font.bold = False
r4_val.font.size = Pt(22)

p5 = tf0.paragraphs[5]
p5.clear()
r5_label = p5.add_run()
r5_label.text = "Team Name - "
r5_label.font.bold = True
r5_label.font.size = Pt(22)
r5_val = p5.add_run()
r5_val.text = "AIZENIX"
r5_val.font.bold = False
r5_val.font.size = Pt(22)

# =========================================================================
# SLIDE 1: Proposed Solution
# =========================================================================
s1 = prs.slides[1]
tf1 = s1.shapes[2].text_frame

p1_2 = tf1.paragraphs[2]
p1_2.clear()
r1_2 = p1_2.add_run()
r1_2.text = "Profiles each candidate/officer (role, experience, past training) & scores competency across 4 core domains to isolate exact skill gaps using deterministic zero-hallucination logic."
r1_2.font.size = Pt(20)

p1_4 = tf1.paragraphs[4]
p1_4.clear()
r1_4 = p1_4.add_run()
r1_4.text = "Recommends personalized iGOT Karmayogi / NSSTA-TPAC courses, tracks completion via iGOT APIs & re-scores after learning — a continuous "
r1_4.font.size = Pt(20)

p1_5 = tf1.paragraphs[5]
p1_5.clear()
r1_5 = p1_5.add_run()
r1_5.text = "   Measure → Diagnose → Learn → Reassess  loop"
r1_5.font.size = Pt(20)
r1_5.font.bold = True

p1_7 = tf1.paragraphs[7]
p1_7.clear()
r1_7_lbl = p1_7.add_run()
r1_7_lbl.text = "Flagship: "
r1_7_lbl.font.bold = True
r1_7_lbl.font.size = Pt(20)
r1_7_txt = p1_7.add_run()
r1_7_txt.text = "trainer uploads a PDF/PPT/DOCX → AI extracts key concepts and generates trainer-reviewed MCQs with rationales; sits ON TOP of iGOT as an intelligence layer, not a replacement."
r1_7_txt.font.size = Pt(20)

# =========================================================================
# SLIDE 2: Technical Approach
# =========================================================================
s2 = prs.slides[2]
tf2 = s2.shapes[2].text_frame

tech_items = [
    ("Frontend: ", "React.js 18, Interactive Dashboards, Recharts (Radar & Bar comparison)"),
    ("Backend: ", "Node.js / Express REST APIs, Deterministic Competency-Scoring Engine"),
    ("AI / NLP: ", "Google Gemini Multimodal APIs, Document Parsing (PDF/DOCX), MCQ Generation"),
    ("RBAC: ", "Role-Based Access Control, Server-Side Trainer Passcode Validation"),
    ("SSO & Security: ", "Single Sign-On, Secure Firebase Authentication, DPDP Act 2023 Compliance"),
    ("API: ", "iGOT REST API Federation Ready, Zero-Infrastructure Serverless Cloud")
]

for idx, (label, val) in enumerate(tech_items):
    p_idx = idx * 2
    if p_idx < len(tf2.paragraphs):
        p = tf2.paragraphs[p_idx]
        p.clear()
        rl = p.add_run()
        rl.text = label
        rl.font.bold = True
        rl.font.size = Pt(20)
        rv = p.add_run()
        rv.text = val
        rv.font.size = Pt(20)

# =========================================================================
# SLIDE 3: Feasibility and Viability
# =========================================================================
s3 = prs.slides[3]
tf3 = s3.shapes[2].text_frame

# p[0], p[1], p[2]
feas = [
    ("MCQ Generation: ", "Can be built using Google Gemini LLM and document parsing technology."),
    ("Competency Scoring: ", "Developed with pure algorithmic logic ensuring 100% mathematical consistency."),
    ("Recommendations: ", "Suggests relevant iGOT micro-courses mapped directly to competency gaps.")
]
for idx, (label, val) in enumerate(feas):
    p = tf3.paragraphs[idx]
    p.clear()
    rl = p.add_run()
    rl.text = label
    rl.font.bold = True
    rl.font.size = Pt(18)
    rv = p.add_run()
    rv.text = val
    rv.font.size = Pt(18)

p_ch_title = tf3.paragraphs[4]
p_ch_title.clear()
r_ch_title = p_ch_title.add_run()
r_ch_title.text = "Challenges:"
r_ch_title.font.bold = True
r_ch_title.font.size = Pt(18)

tf3.paragraphs[5].clear()
r5 = tf3.paragraphs[5].add_run()
r5.text = "• No direct access to iGOT production APIs."
r5.font.size = Pt(17)

tf3.paragraphs[6].clear()
r6 = tf3.paragraphs[6].add_run()
r6.text = "• Candidate and officer data privacy and security concerns."
r6.font.size = Pt(17)

tf3.paragraphs[7].clear()
r7 = tf3.paragraphs[7].add_run()
r7.text = "• Need for an accurate, standardized competency framework."
r7.font.size = Pt(17)

p_sol_title = tf3.paragraphs[9]
p_sol_title.clear()
r_sol_title = p_sol_title.add_run()
r_sol_title.text = "Solution:"
r_sol_title.font.bold = True
r_sol_title.font.size = Pt(18)

tf3.paragraphs[10].clear()
r10 = tf3.paragraphs[10].add_run()
r10.text = "• Start with standardized mock/sandbox iGOT schemas; transition to live gateway via institutional MoU."
r10.font.size = Pt(17)

tf3.paragraphs[11].clear()
r11 = tf3.paragraphs[11].add_run()
r11.text = "• Get official API access gradually through institutional partnership."
r11.font.size = Pt(17)

tf3.paragraphs[12].clear()
r12 = tf3.paragraphs[12].add_run()
r12.text = "• Implement strict DPDP Act 2023 anonymization, server-side auth, and government cloud standards."
r12.font.size = Pt(17)

tf3.paragraphs[13].clear()
r13 = tf3.paragraphs[13].add_run()
r13.text = "• Co-develop the competency framework with NSSTA / MoSPI experts, scaling to RPL frameworks."
r13.font.size = Pt(17)

# =========================================================================
# SLIDE 4: Impact and Benefits (The 5 circular icon shapes!)
# =========================================================================
s4 = prs.slides[4]

# Shape 10
tf10 = s4.shapes[10].text_frame
tf10.paragraphs[0].clear()
r10a = tf10.paragraphs[0].add_run()
r10a.text = "Faster, targeted upskilling for "
r10a.font.bold = True
r10a.font.size = Pt(20)
if len(tf10.paragraphs) > 1:
    tf10.paragraphs[1].clear()
    r10b = tf10.paragraphs[1].add_run()
    r10b.text = "candidates & statistical officers in 15-min micro-learning sprints"
    r10b.font.size = Pt(18)

# Shape 11
tf11 = s4.shapes[11].text_frame
tf11.paragraphs[0].clear()
r11a = tf11.paragraphs[0].add_run()
r11a.text = "Better utilization of existing iGOT content instead of hunting through an unsorted course catalogue."
r11a.font.size = Pt(18)

# Shape 12
tf12 = s4.shapes[12].text_frame
tf12.paragraphs[0].clear()
r12a = tf12.paragraphs[0].add_run()
r12a.text = "Cuts manual course-curation and quiz-writing load on trainers by 80% via automated document ingestion."
r12a.font.size = Pt(18)

# Shape 13
tf13 = s4.shapes[13].text_frame
tf13.paragraphs[0].clear()
r13a = tf13.paragraphs[0].add_run()
r13a.text = "Gives measurable training-effectiveness data and real-time national competency heatmaps."
r13a.font.size = Pt(18)

# Shape 14
tf14 = s4.shapes[14].text_frame
tf14.paragraphs[0].clear()
r14a = tf14.paragraphs[0].add_run()
r14a.text = "Competency model is extensible to other government cadres and nationwide RPL certification."
r14a.font.size = Pt(18)

# =========================================================================
# SLIDE 5: Research, References & Live Prototype
# =========================================================================
s5 = prs.slides[5]
tf5 = s5.shapes[2].text_frame

p5_0 = tf5.paragraphs[0]
p5_0.clear()
r5_0 = p5_0.add_run()
r5_0.text = "iGOT Karmayogi / Mission Karmayogi FRAC framework — igotkarmayogi.gov.in"
r5_0.font.size = Pt(20)

p5_1 = tf5.paragraphs[1]
p5_1.clear()
r5_1 = p5_1.add_run()
r5_1.text = "NSSTA / TPAC training resources — nssta.gov.in"
r5_1.font.size = Pt(20)

p5_2 = tf5.paragraphs[2]
p5_2.clear()
r5_2 = p5_2.add_run()
r5_2.text = "MoSPI official statistics resources — mospi.gov.in  |  DPDP Act 2023 Microdata Standards"
r5_2.font.size = Pt(20)

p5_3 = tf5.paragraphs[3]
p5_3.clear()
r5_3_lbl = p5_3.add_run()
r5_3_lbl.text = "Live Prototype: "
r5_3_lbl.font.bold = True
r5_3_lbl.font.size = Pt(19)
r5_3_url = p5_3.add_run()
r5_3_url.text = "https://statskill-ai-five.vercel.app"
r5_3_url.font.size = Pt(19)
r5_3_url.font.color.rgb = RGBColor(2, 132, 199)

p5_4 = tf5.paragraphs[4]
p5_4.clear()
r5_4_lbl = p5_4.add_run()
r5_4_lbl.text = "GitHub Repository: "
r5_4_lbl.font.bold = True
r5_4_lbl.font.size = Pt(19)
r5_4_url = p5_4.add_run()
r5_4_url.text = "https://github.com/iamwarriorofmyuniverse-coder/statskill-ai"
r5_4_url.font.size = Pt(19)
r5_4_url.font.color.rgb = RGBColor(15, 118, 110)

prs.save(output_path)
print(f"SUCCESS: Saved updated presentation directly into previous template -> {output_path}")

try:
    prs.save(alt_output_path)
    print(f"SUCCESS: Also saved to -> {alt_output_path}")
except Exception as e:
    print(f"Notice: {e}")
