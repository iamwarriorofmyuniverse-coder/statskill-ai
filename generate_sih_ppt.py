import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_sih_presentation():
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Clean Government Navy & Saffron theme
    NAVY = RGBColor(11, 37, 69)       # #0b2545
    OCEAN = RGBColor(2, 132, 199)     # #0284c7
    SAFFRON = RGBColor(234, 88, 12)   # #ea580c
    EMERALD = RGBColor(16, 185, 129)  # #10b981
    DARK = RGBColor(15, 23, 42)       # #0f172a
    MUTED = RGBColor(100, 116, 139)   # #64748b
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)# #f8fafc
    CARD_BG = RGBColor(255, 255, 255)
    BORDER = RGBColor(226, 232, 240)

    def add_header(slide, title, slide_num):
        # Top banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_top = tf.paragraphs[0]
        p_top.text = "SMART INDIA HACKATHON 2026 • AIZENIX (TEAM ID: 176704) • IDEA ID: 21121"
        p_top.font.size = Pt(10)
        p_top.font.bold = True
        p_top.font.color.rgb = SAFFRON

        p_main = tf.add_paragraph()
        p_main.text = title
        p_main.font.size = Pt(22)
        p_main.font.bold = True
        p_main.font.color.rgb = NAVY

        # Slide number bottom right
        num_box = slide.shapes.add_textbox(Inches(11.5), Inches(7.0), Inches(1.0), Inches(0.3))
        tf_n = num_box.text_frame
        p_n = tf_n.paragraphs[0]
        p_n.text = f"{slide_num} | AIZENIX"
        p_n.font.size = Pt(9)
        p_n.font.bold = True
        p_n.font.color.rgb = MUTED
        p_n.alignment = PP_ALIGN.RIGHT

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    
    # Background shape
    bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()

    # Title container
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.0))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "SMART INDIA HACKATHON 2026"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = SAFFRON

    p1 = tf1.add_paragraph()
    p1.text = "StatSkill AI"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(6)

    p2 = tf1.add_paragraph()
    p2.text = "AI-Assisted Skill Assessment & Competency Calibration Platform"
    p2.font.size = Pt(17)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(254, 215, 170)
    p2.space_after = Pt(4)

    p3 = tf1.add_paragraph()
    p3.text = "Adaptive Competency Diagnostic & Automated Capacity Building Platform for Public Sector and RPL Framework"
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(224, 242, 254)
    p3.space_after = Pt(18)

    # Metadata Grid Box
    meta_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.1), Inches(11.333), Inches(2.4))
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = RGBColor(19, 64, 116)
    meta_box.line.color.rgb = RGBColor(56, 189, 248)

    tf_meta = meta_box.text_frame
    tf_meta.word_wrap = True
    tf_meta.margin_left = Inches(0.4)
    tf_meta.margin_top = Inches(0.25)

    pm1 = tf_meta.paragraphs[0]
    pm1.text = "• Theme: Smart Education   |   • Category: Software   |   • Target Ecosystem: Mission Karmayogi (iGOT) & RPL"
    pm1.font.size = Pt(12)
    pm1.font.bold = True
    pm1.font.color.rgb = WHITE
    pm1.space_after = Pt(6)

    pm2 = tf_meta.add_paragraph()
    pm2.text = "• Team Name: AIZENIX       |   • Team ID: 176704      |   • Idea / PS ID: 21121"
    pm2.font.size = Pt(13)
    pm2.font.bold = True
    pm2.font.color.rgb = RGBColor(254, 215, 170)
    pm2.space_after = Pt(6)

    pm3 = tf_meta.add_paragraph()
    pm3.text = "• Live Prototype: https://statskill-ai-five.vercel.app   |   • GitHub: https://github.com/iamwarriorofmyuniverse-coder/statskill-ai"
    pm3.font.size = Pt(11)
    pm3.font.color.rgb = RGBColor(56, 189, 248)

    # =========================================================================
    # SLIDE 2: Proposed Solution
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Proposed Solution: Competency Layer for iGOT Karmayogi", 2)

    cards = [
        ("1. Role-Calibrated Capability Profiling", 
         "Profiles each officer (role, seniority, past training history) and scores competencies across 4 domains (Statistical Methods, Technical Analytics, Governance/DPDP 2023, and Behavioural Skills) to isolate exact skill gaps with zero-hallucination deterministic JavaScript logic.",
         NAVY),
        ("2. Closed-Loop Remediation Cycle",
         "Recommends targeted iGOT / NSSTA-TPAC courses, tracks completion milestones via iGOT APIs, and re-scores after learning in a continuous:\n👉 Measure → Diagnose → Learn → Reassess loop.",
         OCEAN),
        ("3. AI-Powered MCQ Generator & Document Intelligence",
         "Flagship Trainer Engine: Training managers upload reference circulars, manuals (PDF/DOCX/PPT) → AI extracts key concepts and generates trainer-reviewed MCQs with explanations.\nSits ON TOP of iGOT as an intelligence layer without replacing existing LMS.",
         SAFFRON)
    ]

    for i, (head, body, color) in enumerate(cards):
        top_pos = Inches(1.5 + i * 1.8)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.733), Inches(1.6))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.3)

        p_h = tf.paragraphs[0]
        p_h.text = head
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = color
        p_h.space_after = Pt(4)

        p_b = tf.add_paragraph()
        p_b.text = body
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = DARK

    # =========================================================================
    # SLIDE 3: Technical Approach
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Technical Approach & Architecture", 3)

    tech_pillars = [
        ("Frontend & UI/UX", "• React.js 18 + Vite & Tailwind CSS\n• Interactive Recharts (Radar + Bar comparison)\n• Fully responsive (Mobile, Tablet, Desktop)\n• WCAG AAA High-Contrast Dark/Light themes"),
        ("Backend & Cloud", "• Node.js & Express REST APIs\n• Deterministic JS Assessment Engine\n• Firebase Cloud Firestore for real-time sync\n• Vercel Serverless auto-scaling deployment"),
        ("AI / NLP & Voice", "• Google Gemini Multimodal APIs\n• Automated PDF/DOCX Parsing & MCQ extraction\n• Bilingual Web Speech API (English & Hindi)\n• Hands-free cadre voice tutoring"),
        ("Security, RBAC & SSO", "• Role-Based Access (Officer, Trainer, Admin)\n• Secure Firebase Tokenized Authentication\n• DPDP Act 2023 microdata privacy standards\n• iGOT Karmayogi REST API Federation ready")
    ]

    for i, (title, content) in enumerate(tech_pillars):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 6.0)
        top = Inches(1.5 + row * 2.7)

        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.733), Inches(2.5))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = BORDER
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.25)

        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = NAVY
        pt.space_after = Pt(8)

        pc = tf.add_paragraph()
        pc.text = content
        pc.font.size = Pt(11)
        pc.font.color.rgb = DARK

    # =========================================================================
    # SLIDE 4: Feasibility and Viability
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Feasibility, Viability & Mitigation Strategy", 4)

    # Left: Feasibility pillars
    f_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.3))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = CARD_BG
    f_box.line.color.rgb = EMERALD
    f_box.line.width = Pt(1.5)

    tf_f = f_box.text_frame
    tf_f.word_wrap = True
    tf_f.margin_left = Inches(0.3)
    tf_f.margin_top = Inches(0.3)

    pf0 = tf_f.paragraphs[0]
    pf0.text = "Technical Feasibility"
    pf0.font.size = Pt(14)
    pf0.font.bold = True
    pf0.font.color.rgb = EMERALD
    pf0.space_after = Pt(10)

    feas_points = [
        "• MCQ Generation: Production-ready LLM pipelines with document parsing and prompt calibration.",
        "• Deterministic Scoring: Pure JavaScript algorithmic engine guaranteeing 100% mathematical consistency without hallucinations.",
        "• Course Recommendations: Automated mapping of deficiency vectors (Gap = Target - Current) to iGOT competency modules.",
        "• Low-Bandwidth Operation: Lightweight client assets (<500KB gzipped) optimized for rural field investigators."
    ]
    for p in feas_points:
        p_item = tf_f.add_paragraph()
        p_item.text = p
        p_item.font.size = Pt(11)
        p_item.font.color.rgb = DARK
        p_item.space_after = Pt(8)

    # Right: Challenges & Solutions
    c_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.733), Inches(5.3))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = CARD_BG
    c_box.line.color.rgb = SAFFRON
    c_box.line.width = Pt(1.5)

    tf_c = c_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = Inches(0.3)
    tf_c.margin_top = Inches(0.3)

    pc0 = tf_c.paragraphs[0]
    pc0.text = "Challenges & Institutional Solutions"
    pc0.font.size = Pt(14)
    pc0.font.bold = True
    pc0.font.color.rgb = SAFFRON
    pc0.space_after = Pt(10)

    mitigations = [
        ("Challenge: Direct iGOT API Access", "Solution: Start with mock/sandbox iGOT API schemas; transition to live government gateway through institutional MoU."),
        ("Challenge: Data Privacy & Security", "Solution: Strict DPDP Act 2023 data anonymization, zero plain-text storage of officer scores, and tokenized access control."),
        ("Challenge: Cadre Framework Alignment", "Solution: Designed around NSSTA / MoSPI competency standards (NSSO, ASI, PLFS) with extensible taxonomy for other civil cadres.")
    ]
    for ch, sol in mitigations:
        p_ch = tf_c.add_paragraph()
        p_ch.text = f"❌ {ch}"
        p_ch.font.size = Pt(11)
        p_ch.font.bold = True
        p_ch.font.color.rgb = RGBColor(185, 28, 28)

        p_sol = tf_c.add_paragraph()
        p_sol.text = f"✅ {sol}"
        p_sol.font.size = Pt(10.5)
        p_sol.font.color.rgb = DARK
        p_sol.space_after = Pt(8)

    # =========================================================================
    # SLIDE 5: Impact and Benefits
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Measurable Impact & National Benefits", 5)

    benefits = [
        ("Faster, Targeted Upskilling for Statistical Cadres", "Officers address specific weaknesses (e.g. Sampling, DPDP compliance, Python) in 15-minute micro-learning sprints rather than 40-hour generic courses.", OCEAN),
        ("Optimal Utilization of Existing iGOT Assets", "Increases course completion rates by precision-routing learners directly to relevant modules instead of browsing unindexed catalogues.", NAVY),
        ("Cuts Trainer Workload by 80%", "Automated document ingestion generates verified, curriculum-aligned MCQ banks in seconds, eliminating manual test creation.", SAFFRON),
        ("Actionable Cadre Competency Telemetry", "Provides ministry leadership and NSSTA with real-time heatmaps of national statistical readiness and skill gaps.", EMERALD),
        ("Universal & Extensible Across Ministries", "While piloted on MoSPI statistical cadres, the diagnostic engine readily scales to State DES, Finance, and Public Administration cadres.", RGBColor(124, 58, 237))
    ]

    for i, (title, desc, color) in enumerate(benefits):
        top_pos = Inches(1.5 + i * 1.08)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.733), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.12)
        tf.margin_right = Inches(0.25)

        pt = tf.paragraphs[0]
        pt.text = f"✔ {title}"
        pt.font.size = Pt(11.5)
        pt.font.bold = True
        pt.font.color.rgb = color

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(10)
        pd.font.color.rgb = DARK

    # =========================================================================
    # SLIDE 6: Research, References & Live Prototype
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Research, References & Live Working Prototype", 6)

    # Left Box: References
    r_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    r_box.fill.solid()
    r_box.fill.fore_color.rgb = CARD_BG
    r_box.line.color.rgb = BORDER
    r_box.line.width = Pt(1.5)

    tf_r = r_box.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_top = Inches(0.3)

    pr0 = tf_r.paragraphs[0]
    pr0.text = "Institutional Research & Standards"
    pr0.font.size = Pt(14)
    pr0.font.bold = True
    pr0.font.color.rgb = NAVY
    pr0.space_after = Pt(12)

    refs = [
        ("iGOT Karmayogi / Mission Karmayogi", "Framework of Roles, Activities and Competencies (FRAC) — igotkarmayogi.gov.in"),
        ("NSSTA / TPAC Training Academy", "National Statistical Systems Training Academy Competency Standards — nssta.gov.in"),
        ("Ministry of Statistics & PI (MoSPI)", "Official Statistical System Standards, NSSO / ASI / PLFS Guidelines — mospi.gov.in"),
        ("Digital Personal Data Protection (DPDP)", "DPDP Act 2023 microdata privacy and anonymization specifications — meity.gov.in")
    ]
    for h, d in refs:
        ph = tf_r.add_paragraph()
        ph.text = f"• {h}"
        ph.font.size = Pt(11)
        ph.font.bold = True
        ph.font.color.rgb = OCEAN

        pd = tf_r.add_paragraph()
        pd.text = d
        pd.font.size = Pt(10)
        pd.font.color.rgb = DARK
        pd.space_after = Pt(8)

    # Right Box: Prominent Live Prototype & GitHub Banner
    p_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.733), Inches(5.3))
    p_box.fill.solid()
    p_box.fill.fore_color.rgb = RGBColor(240, 249, 255) # Light sky blue
    p_box.line.color.rgb = OCEAN
    p_box.line.width = Pt(2.0)

    tf_p = p_box.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = Inches(0.35)
    tf_p.margin_top = Inches(0.35)

    pp0 = tf_p.paragraphs[0]
    pp0.text = "🚀 Live Working Prototype & Source Code"
    pp0.font.size = Pt(15)
    pp0.font.bold = True
    pp0.font.color.rgb = NAVY
    pp0.space_after = Pt(14)

    # Live Vercel Box
    pv1 = tf_p.add_paragraph()
    pv1.text = "🌐 Live Production URL:"
    pv1.font.size = Pt(11)
    pv1.font.bold = True
    pv1.font.color.rgb = DARK

    pv2 = tf_p.add_paragraph()
    pv2.text = "https://statskill-ai-five.vercel.app"
    pv2.font.size = Pt(12)
    pv2.font.bold = True
    pv2.font.color.rgb = RGBColor(2, 132, 199)
    pv2.space_after = Pt(12)

    # GitHub Box
    pg1 = tf_p.add_paragraph()
    pg1.text = "💻 GitHub Repository (Open Source):"
    pg1.font.size = Pt(11)
    pg1.font.bold = True
    pg1.font.color.rgb = DARK

    pg2 = tf_p.add_paragraph()
    pg2.text = "https://github.com/iamwarriorofmyuniverse-coder/statskill-ai"
    pg2.font.size = Pt(10.5)
    pg2.font.bold = True
    pg2.font.color.rgb = RGBColor(15, 118, 110)
    pg2.space_after = Pt(14)

    # Highlights
    p_feat = tf_p.add_paragraph()
    p_feat.text = "Platform Highlights Ready for Live Demo:"
    p_feat.font.size = Pt(11)
    p_feat.font.bold = True
    p_feat.font.color.rgb = SAFFRON
    p_feat.space_after = Pt(4)

    features = [
        "✔ 22-Question Deterministic Cadre Diagnostic Engine",
        "✔ Real-Time Radar Competency Calibration & APAR Export",
        "✔ Multimodal Bilingual AI Voice Assistant (English + Hindi)",
        "✔ Interactive iGOT Micro-Learning Console & Chapter Quizzes",
        "✔ 1-Page Printable MoSPI / NSSTA Verified Cadre Certificate"
    ]
    for f in features:
        pf = tf_p.add_paragraph()
        pf.text = f
        pf.font.size = Pt(10)
        pf.font.color.rgb = DARK
        pf.space_after = Pt(2)

    output_path = os.path.join(os.getcwd(), "SIH_PPT_AIZENIX_176704.pptx")
    prs.save(output_path)
    print(f"SUCCESS: Created {output_path}")

if __name__ == "__main__":
    create_sih_presentation()
